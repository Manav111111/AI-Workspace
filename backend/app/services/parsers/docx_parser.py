import io
from typing import List
import docx
from app.core.exceptions import ValidationException
from app.services.parsers.base import DocumentParser, ParsedSection


class DOCXParser(DocumentParser):
    def can_parse(self, file_type: str) -> bool:
        return file_type.lower() in [
            "docx",
            ".docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ]

    def parse(self, content: bytes, original_filename: str) -> List[ParsedSection]:
        try:
            doc = docx.Document(io.BytesIO(content))
        except Exception as e:
            raise ValidationException(f"Failed to read DOCX document: {str(e)}")

        sections: List[ParsedSection] = []
        current_section_title = "Introduction"
        current_paragraphs: List[str] = []

        # Process paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            # Heading detection
            if para.style.name.startswith("Heading"):
                if current_paragraphs:
                    sections.append(
                        ParsedSection(
                            content="\n".join(current_paragraphs),
                            metadata={
                                "section": current_section_title,
                                "source": original_filename,
                                "file_type": "docx",
                            },
                        )
                    )
                    current_paragraphs = []
                current_section_title = text
            else:
                current_paragraphs.append(text)

        if current_paragraphs:
            sections.append(
                ParsedSection(
                    content="\n".join(current_paragraphs),
                    metadata={
                        "section": current_section_title,
                        "source": original_filename,
                        "file_type": "docx",
                    },
                )
            )

        # Process tables
        for table_idx, table in enumerate(doc.tables):
            table_lines: List[str] = []
            headers: List[str] = []
            for row_idx, row in enumerate(table.rows):
                row_cells = [cell.text.strip() for cell in row.cells]
                if row_idx == 0:
                    headers = row_cells
                    table_lines.append(" | ".join(row_cells))
                else:
                    if headers and len(headers) == len(row_cells):
                        formatted_row = " | ".join(
                            f"{headers[i]}: {row_cells[i]}" for i in range(len(headers))
                        )
                        table_lines.append(formatted_row)
                    else:
                        table_lines.append(" | ".join(row_cells))

            if table_lines:
                sections.append(
                    ParsedSection(
                        content="\n".join(table_lines),
                        metadata={
                            "section": f"Table {table_idx + 1}",
                            "source": original_filename,
                            "file_type": "docx",
                            "is_table": True,
                        },
                    )
                )

        if not sections:
            raise ValidationException("DOCX contains no extractable text")

        return sections
