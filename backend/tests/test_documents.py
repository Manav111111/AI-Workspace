import io
import docx
import fitz
import pytest
from httpx import AsyncClient


def create_sample_pdf() -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "AI Employee Platform Return Policy.\nCustomers may return products within 30 days.")
    page2 = doc.new_page()
    page2.insert_text((50, 72), "Page 2: Warranty Information.\nAll products carry a 1-year warranty.")
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_sample_docx() -> bytes:
    doc = docx.Document()
    doc.add_heading("Company Handbook", level=1)
    doc.add_paragraph("Welcome to the AI Employee Platform. We build multi-tenant AI agents.")
    doc.add_heading("Working Hours", level=2)
    doc.add_paragraph("Standard working hours are flexible between 9am and 5pm.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Role"
    table.rows[0].cells[1].text = "Tier"
    table.rows[1].cells[0].text = "Support Agent"
    table.rows[1].cells[1].text = "Level 1"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_document_uploads_and_formats(client: AsyncClient):
    # Setup company and KB
    signup = await client.post(
        "/api/v1/auth/signup",
        json={
            "email": "docs_admin@company.com",
            "password": "Password123!",
            "full_name": "Docs Admin",
            "company_name": "DocuCorp",
        },
    )
    token = signup.json()["token"]["access_token"]
    company_id = signup.json()["company"]["id"]
    headers = {"Authorization": f"Bearer {token}", "X-Company-ID": company_id}

    kb_res = await client.post(
        "/api/v1/knowledge-bases/",
        json={"name": "Primary Knowledge Base"},
        headers=headers,
    )
    kb_id = kb_res.json()["id"]

    # 1. Upload PDF
    pdf_bytes = create_sample_pdf()
    res_pdf = await client.post(
        f"/api/v1/knowledge-bases/{kb_id}/documents",
        files={"file": ("policy.pdf", pdf_bytes, "application/pdf")},
        headers=headers,
    )
    assert res_pdf.status_code == 201
    doc_pdf = res_pdf.json()
    assert doc_pdf["file_type"] == "pdf"
    assert doc_pdf["status"] == "PROCESSED"
    assert doc_pdf["document_metadata"]["total_chunks"] >= 2
    pdf_id = doc_pdf["id"]

    # 2. Upload DOCX
    docx_bytes = create_sample_docx()
    res_docx = await client.post(
        f"/api/v1/knowledge-bases/{kb_id}/documents",
        files={"file": ("handbook.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        headers=headers,
    )
    assert res_docx.status_code == 201
    doc_docx = res_docx.json()
    assert doc_docx["file_type"] == "docx"
    assert doc_docx["status"] == "PROCESSED"

    # 3. Upload Markdown
    md_content = b"# Architecture Blueprint\n\nThe platform is a modular monolith.\n\n## Multi-Tenancy\nTenant isolation is enforced strictly."
    res_md = await client.post(
        f"/api/v1/knowledge-bases/{kb_id}/documents",
        files={"file": ("architecture.md", md_content, "text/markdown")},
        headers=headers,
    )
    assert res_md.status_code == 201
    doc_md = res_md.json()
    assert doc_md["file_type"] == "md"
    assert doc_md["status"] == "PROCESSED"

    # 4. Upload TXT
    txt_content = b"Simple text note.\nSecond paragraph with details."
    res_txt = await client.post(
        f"/api/v1/knowledge-bases/{kb_id}/documents",
        files={"file": ("notes.txt", txt_content, "text/plain")},
        headers=headers,
    )
    assert res_txt.status_code == 201
    doc_txt = res_txt.json()
    assert doc_txt["file_type"] == "txt"
    assert doc_txt["status"] == "PROCESSED"

    # 5. Upload CSV
    csv_content = b"Product,Price,Category\nPro Plan,$49,Subscription\nEnterprise Plan,$499,Enterprise"
    res_csv = await client.post(
        f"/api/v1/knowledge-bases/{kb_id}/documents",
        files={"file": ("pricing.csv", csv_content, "text/csv")},
        headers=headers,
    )
    assert res_csv.status_code == 201
    doc_csv = res_csv.json()
    assert doc_csv["file_type"] == "csv"
    assert doc_csv["status"] == "PROCESSED"

    # 6. Test Chunks Retrieval
    chunks_res = await client.get(f"/api/v1/documents/{pdf_id}/chunks", headers=headers)
    assert chunks_res.status_code == 200
    chunks = chunks_res.json()
    assert len(chunks) >= 2
    assert "page" in chunks[0]["chunk_metadata"]
    assert chunks[0]["company_id"] == company_id

    # 7. Validation: Unsupported File Format
    res_bad = await client.post(
        f"/api/v1/knowledge-bases/{kb_id}/documents",
        files={"file": ("malware.exe", b"binary content", "application/octet-stream")},
        headers=headers,
    )
    assert res_bad.status_code == 422
    assert res_bad.json()["error"]["code"] == "VALIDATION_ERROR"

    # 8. Validation: Empty File
    res_empty = await client.post(
        f"/api/v1/knowledge-bases/{kb_id}/documents",
        files={"file": ("empty.txt", b"", "text/plain")},
        headers=headers,
    )
    assert res_empty.status_code == 422
    assert res_empty.json()["error"]["code"] == "VALIDATION_ERROR"

    # 9. Delete Document
    del_res = await client.delete(f"/api/v1/documents/{pdf_id}", headers=headers)
    assert del_res.status_code == 204

    # Verify deleted
    verify_doc = await client.get(f"/api/v1/documents/{pdf_id}", headers=headers)
    assert verify_doc.status_code == 404
